import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from document_parser import MAX_DOCUMENT_BYTES, parse_document


def pdf_bytes(text=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = writer._add_object(DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }))
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 30 200 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def word_bytes():
    document = Document()
    document.add_paragraph("galacticworddeadline is in the Word document.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Word table"
    table.cell(0, 1).text = "wordtablemarker"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def ppt_bytes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text = "orbitalpptdeadline is in the slide."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "PPT table"
    table.cell(0, 1).text = "ppttablemarker"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class DocumentParserTests(unittest.TestCase):
    def test_three_formats_extract_text_and_sources(self):
        cases = [
            ("sample.pdf", pdf_bytes("nebularpdfdeadline"), "nebularpdfdeadline", "page 1"),
            ("sample.docx", word_bytes(), "wordtablemarker", "sample.docx"),
            ("sample.pptx", ppt_bytes(), "ppttablemarker", "slide 1"),
        ]
        for name, content, phrase, source in cases:
            with self.subTest(name=name):
                chunks = parse_document(name, content)
                self.assertTrue(any(phrase in chunk["text"] for chunk in chunks))
                self.assertTrue(any(source in chunk["source"] for chunk in chunks))
                self.assertTrue(all(chunk["id"].startswith("DOC_") for chunk in chunks))

    def test_empty_pdf_unsupported_format_and_size_limit(self):
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_document("scan.pdf", pdf_bytes())
        with self.assertRaisesRegex(ValueError, "PDF, DOCX and PPTX"):
            parse_document("old.doc", b"data")
        with self.assertRaisesRegex(ValueError, "10 MB"):
            parse_document("large.pdf", b"x" * (MAX_DOCUMENT_BYTES + 1))


class DocumentAppTests(unittest.TestCase):
    def start_app(self):
        app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
        self.assertFalse(app.exception)
        uploader = next((x for x in app.get("file_uploader") if "文档" in getattr(x, "label", "")), None)
        if uploader is None or not hasattr(uploader, "set_value"):
            self.skipTest("This Streamlit version cannot simulate file uploads")
        return app, uploader

    def test_uploaded_documents_are_searchable_and_faq_evaluation_stays_separate(self):
        app, uploader = self.start_app()
        uploader.set_value([
            ("sample.pdf", pdf_bytes("nebularpdfdeadline"), "application/pdf"),
            ("sample.docx", word_bytes(), "application/octet-stream"),
            ("sample.pptx", ppt_bytes(), "application/octet-stream"),
        ]).run()
        self.assertFalse(app.exception)
        for query, expected in [
            ("nebularpdfdeadline", "sample.pdf (page 1)"),
            ("galacticworddeadline", "sample.docx"),
            ("orbitalpptdeadline", "sample.pptx (slide 1)"),
        ]:
            with self.subTest(query=query):
                app.chat_input[0].set_value(query).run()
                self.assertFalse(app.exception)
                self.assertIn(expected, app.chat_message[-1].markdown[0].value)
        app.chat_input[0].set_value("如何申请作业延期？").run()
        self.assertIn("演示政策", app.chat_message[-1].markdown[0].value)
        next(button for button in app.button if button.label == "运行评估").click().run()
        self.assertFalse(app.exception)
        _, predictions, _ = app.session_state["evaluation"]
        self.assertFalse(predictions["retrieved"].str.contains("DOC_").any())

    def test_invalid_document_does_not_break_faq(self):
        app, uploader = self.start_app()
        uploader.upload("bad.pdf", b"not a PDF").run()
        self.assertFalse(app.exception)
        self.assertTrue(any("bad.pdf" in warning.value for warning in app.warning))
        app.chat_input[0].set_value("如何申请作业延期？").run()
        self.assertIn("演示政策", app.chat_message[-1].markdown[0].value)

    def test_rag_accepts_document_citation(self):
        app, uploader = self.start_app()
        uploader.upload("sample.pdf", pdf_bytes("nebularpdfdeadline")).run()
        document_id = next(x["id"] for x in app.session_state["engine"].data if x["kind"] == "document")
        app.radio[0].set_value("RAG").run()
        with patch("requests.post") as post:
            post.return_value.json.return_value = {
                "message": {"content": f"The PDF describes the deadline. [{document_id}]"}
            }
            app.chat_input[0].set_value("nebularpdfdeadline").run()
        self.assertFalse(app.exception)
        self.assertIn(f"[{document_id}]", app.chat_message[-1].markdown[0].value)
        post.assert_called_once()


if __name__ == "__main__":
    unittest.main()
