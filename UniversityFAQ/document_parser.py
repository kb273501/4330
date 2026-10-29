"""Extract searchable text from uploaded documents without writing them to disk."""

import hashlib
import re
from io import BytesIO
from pathlib import Path


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_DOCUMENT_CHARS = 2_000_000
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 120


def _chunks(text):
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = max(
                text.rfind("\n", start + CHUNK_SIZE // 2, end),
                text.rfind(" ", start + CHUNK_SIZE // 2, end),
            )
            if boundary != -1:
                end = boundary + 1
        chunk = text[start:end].strip()
        if chunk:
            yield chunk
        if end == len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)


def _pdf_sections(content):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("Password-protected PDFs are not supported / 不支持加密 PDF")
    for number, page in enumerate(reader.pages, 1):
        yield f"page {number}", page.extract_text() or ""


def _word_sections(content):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(content))
    lines = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        elif isinstance(block, Table):
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
    yield "", "\n".join(lines)


def _slide_text(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_slide_text(shape.shapes))
        elif shape.has_text_frame:
            lines.append(shape.text_frame.text)
        elif shape.has_table:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return lines


def _ppt_sections(content):
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    for number, slide in enumerate(presentation.slides, 1):
        yield f"slide {number}", "\n".join(_slide_text(slide.shapes))


def parse_document(name, content):
    """Return document chunks with stable IDs and source locations."""
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("Maximum file size is 10 MB / 单个文件最多 10 MB")

    suffix = Path(name).suffix.lower()
    extractors = {".pdf": _pdf_sections, ".docx": _word_sections, ".pptx": _ppt_sections}
    if suffix not in extractors:
        raise ValueError("Only PDF, DOCX and PPTX are supported / 仅支持 PDF、DOCX 和 PPTX")

    digest = hashlib.sha256(name.encode("utf-8") + b"\0" + content).hexdigest()[:16]
    chunks = []
    total_chars = 0
    for location, text in extractors[suffix](content):
        total_chars += len(text)
        if total_chars > MAX_DOCUMENT_CHARS:
            raise ValueError("Document text is too long / 文档文字过长")
        source = f"{name} ({location})" if location else name
        for chunk in _chunks(text):
            chunks.append({"id": f"DOC_{digest}_{len(chunks) + 1}", "text": chunk, "source": source})
    if not chunks:
        raise ValueError("No selectable text found (scanned images need OCR) / 未找到可提取文字（扫描图片需要 OCR）")
    return chunks
